import hashlib
import io
import json
import sqlite3
from openpyxl import Workbook
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi import HTTPException
from starlette.datastructures import UploadFile
from starlette.requests import Request
from starlette.responses import Response

from app import main
from app.config import get_settings
from app.schemas import DatabaseSelectRequest, ModelQuery
from app.workflow import run_query


def request_for(session):
    return Request({"type":"http","headers":[(b"x-querypilot-session",session.encode())],"method":"GET","path":"/api/databases","query_string":b"","server":("test",80),"client":("test",1),"scheme":"http","http_version":"1.1"})


def make_database(path, table, value):
    con = sqlite3.connect(path)
    con.executescript(f"CREATE TABLE parent(id INTEGER PRIMARY KEY, name TEXT NOT NULL); CREATE TABLE {table}(id INTEGER PRIMARY KEY, parent_id INTEGER REFERENCES parent(id), amount REAL); INSERT INTO parent VALUES(1,'{value}'); INSERT INTO {table} VALUES(1,1,12.5);")
    con.commit(); con.close()
    return path.read_bytes()


def upload(session, filename, content):
    req = request_for(session); response = Response()
    file = UploadFile(filename=filename, file=io.BytesIO(content))
    return main.upload(req, response, file)


@pytest.fixture
def setup(tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOADS_DIR", str(tmp_path / "uploads"))
    get_settings().max_upload_bytes = 1024 * 1024
    get_settings().max_databases_per_session = 5
    return tmp_path


def test_upload_schema_session_isolation_and_switching(setup):
    content = make_database(setup / "a.sqlite", "sales", "A")
    a, b = "a" * 32, "b" * 32
    body = upload(a, "sample.sqlite", content)
    assert body["schema"]["tables"]["sales"]["primary_keys"] == ["id"]
    assert body["schema"]["tables"]["sales"]["foreign_keys"] == [{"column":"parent_id","referenced_table":"parent","referenced_column":"id"}]
    schema = main.schema(request_for(a), Response())
    assert set(schema["tables"]) == {"parent", "sales"}
    assert set(main.schema(request_for(b), Response())["tables"]) == {"customers", "products", "orders", "order_items"}
    with pytest.raises(HTTPException) as error:
        main.activate(DatabaseSelectRequest(database_id=body["id"]), request_for(b), Response())
    assert error.value.status_code == 400
    assert main.activate(DatabaseSelectRequest(database_id="sample"), request_for(a), Response())["is_sample"] is True
    assert main.activate(DatabaseSelectRequest(database_id=body["id"]), request_for(a), Response())["name"] == "sample.sqlite"


def test_invalid_and_size_limited_upload_preserve_active_selection(setup):
    session = "c" * 32
    content = make_database(setup / "ok.db", "things", "ok")
    active = upload(session, "ok.db", content)["id"]
    with pytest.raises(HTTPException):
        upload(session, "broken.db", b"not sqlite")
    assert main.active(request_for(session), Response())["id"] == active
    get_settings().max_upload_bytes = 10
    with pytest.raises(HTTPException):
        upload(session, "large.sqlite", content)
    get_settings().max_upload_bytes = 1024 * 1024


def test_uploaded_query_validation_and_database_unchanged(setup):
    session = "d" * 32
    source = setup / "query.sqlite3"
    original = make_database(source, "sales", "buyer")
    body = upload(session, source.name, original)
    result = run_query("Show sales", lambda *args: ModelQuery(sql="SELECT s.amount, p.name FROM sales s JOIN parent p ON p.id=s.parent_id", explanation="show sales"), session_id=session)
    assert result["status"] == "success"
    assert result["rows"] == [{"amount":12.5,"name":"buyer"}]
    assert run_query("bad", lambda *args: ModelQuery(sql="DELETE FROM sales", explanation="write"), session_id=session)["status"] == "refused"
    assert run_query("bad table", lambda *args: ModelQuery(sql="SELECT * FROM customers", explanation="read"), session_id=session)["status"] == "refused"
    assert run_query("bad column", lambda *args: ModelQuery(sql="SELECT secret FROM sales", explanation="read"), session_id=session)["status"] == "refused"
    directory = setup / "uploads" / hashlib.sha256(session.encode()).hexdigest()
    entry = json.loads((directory / "registry.json").read_text())["databases"][body["id"]]
    assert (directory / entry["file"]).read_bytes() == original


def test_wal_format_upload_is_rejected_and_sample_remains_active(setup):
    session = "e" * 32
    content = bytearray(make_database(setup / "wal.sqlite", "items", "x"))
    content[18] = 2
    with pytest.raises(HTTPException):
        upload(session, "wal.sqlite", bytes(content))
    assert main.active(request_for(session), Response())["is_sample"] is True


def test_ordinary_views_are_accepted_but_generated_columns_are_rejected(setup):
    session = "7" * 32
    view_db = sqlite3.connect(setup / "view.db")
    view_db.executescript("CREATE TABLE base(id INTEGER PRIMARY KEY); CREATE VIEW exposed AS SELECT id FROM base;")
    view_db.close()
    accepted = upload(session, "view.db", (setup / "view.db").read_bytes())
    assert "exposed" in accepted["schema"]["views"]
    active_id = accepted["id"]
    generated_db = sqlite3.connect(setup / "generated.db")
    generated_db.execute("CREATE TABLE calc(a INTEGER, b INTEGER GENERATED ALWAYS AS (a + 1) STORED)")
    generated_db.close()
    with pytest.raises(HTTPException) as error:
        upload(session, "generated.db", (setup / "generated.db").read_bytes())
    assert "generated or hidden" in error.value.detail
    assert main.active(request_for(session), Response())["id"] == active_id


def test_employee_department_view_upload_query_and_safety(setup):
    session = "8" * 32
    source = setup / "employees.sqlite"
    con = sqlite3.connect(source)
    con.executescript("""
        CREATE TABLE employee(employee_id INTEGER PRIMARY KEY, full_name TEXT NOT NULL, job_role TEXT NOT NULL, department_id INTEGER REFERENCES department(department_id), private_note TEXT);
        CREATE TABLE department(department_id INTEGER PRIMARY KEY, department_name TEXT NOT NULL);
        INSERT INTO department VALUES (1, 'Engineering');
        INSERT INTO employee VALUES (10, 'Avery Chen', 'Analyst', 1, 'confidential');
        CREATE VIEW Join_employee_and_department_tables AS
          SELECT e.full_name, e.job_role, d.department_name
          FROM employee AS e JOIN department AS d ON d.department_id = e.department_id;
    """)
    con.commit(); con.close()

    body = upload(session, source.name, source.read_bytes())
    view = body["schema"]["views"]["Join_employee_and_department_tables"]
    assert set(view["columns"]) == {"full_name", "job_role", "department_name"}
    assert set(view["depends_on"]) == {"employee", "department"}
    active_schema = main.active(request_for(session), Response())["schema"]
    assert "Join_employee_and_department_tables" in active_schema["views"]

    generated = lambda sql: ModelQuery(sql=sql, explanation="Read employee directory.")
    result = run_query(
        "List employees in Engineering",
        lambda *args: generated("SELECT full_name, job_role, department_name FROM Join_employee_and_department_tables WHERE department_name = 'Engineering'"),
        session_id=session,
    )
    assert result["status"] == "success"
    assert result["rows"] == [{"full_name":"Avery Chen","job_role":"Analyst","department_name":"Engineering"}]

    for sql in (
        "SELECT * FROM unknown_view",
        "SELECT private_note FROM Join_employee_and_department_tables",
        "SELECT * FROM sqlite_master",
        "DELETE FROM employee",
        "SELECT full_name FROM Join_employee_and_department_tables; SELECT 1",
    ):
        rejected = run_query("unsafe request", lambda *args, sql=sql: generated(sql), session_id=session)
        assert rejected["status"] == "refused", (sql, rejected)


def test_view_with_unauthorized_dependency_is_rejected_at_upload(setup):
    session = "9" * 32
    source = setup / "unsafe_view.sqlite"
    con = sqlite3.connect(source)
    con.executescript("CREATE TABLE employee(name TEXT); CREATE VIEW bad_view AS SELECT name FROM sqlite_master;")
    con.close()
    with pytest.raises(HTTPException) as error:
        upload(session, source.name, source.read_bytes())
    assert "not available" in error.value.detail.lower()
    assert main.active(request_for(session), Response())["is_sample"] is True


def test_view_with_unsafe_function_is_rejected_without_invoking_it(setup):
    session = "6" * 32
    source = setup / "unsafe_function.sqlite"
    con = sqlite3.connect(source)
    con.executescript("CREATE TABLE employee(name TEXT); CREATE VIEW unsafe_view AS SELECT load_extension('/tmp/not-an-extension') AS name FROM employee;")
    con.close()
    with pytest.raises(HTTPException) as error:
        upload(session, source.name, source.read_bytes())
    assert "function" in error.value.detail.lower()
    assert main.active(request_for(session), Response())["is_sample"] is True


def test_csv_tsv_uploads_convert_to_sqlite_and_query_same_active_schema(setup):
    session = "1" * 32
    csv_data = "\ufeffid,name,amount,created\n1,Avery,12.5,2026-09-01\n2,Jordan,8,2026-09-30\n".encode()
    result = upload(session, "sales.csv", csv_data)
    assert result["source_type"] == "csv"
    assert result["original_filename"] == "sales.csv"
    assert result["schema"]["tables"]["sales"]["columns"]["id"]["type"] == "INTEGER"
    assert result["schema"]["tables"]["sales"]["columns"]["amount"]["type"] == "REAL"
    listed = main.databases(request_for(session), Response())
    assert listed["databases"][-1]["source_type"] == "csv"
    assert listed["max_upload_bytes"] == get_settings().max_upload_bytes
    assert listed["max_import_cells"] == get_settings().max_import_cells
    assert listed["max_import_rows"] == get_settings().max_import_rows
    query = run_query("show imported rows", lambda *args: ModelQuery(sql="SELECT id, name, amount FROM sales WHERE created >= '2026-09-01' AND created < '2026-10-01' ORDER BY amount DESC LIMIT 1", explanation="show imported rows"), session_id=session)
    assert query["status"] == "success"
    assert query["rows"] == [{"id":1,"name":"Avery","amount":12.5}]

    tsv = upload(session, "people.tsv", "person\tactive\nLee\t1\n".encode())
    assert tsv["schema"]["tables"]["people"]["columns"]["person"]["type"] == "TEXT"
    assert tsv["schema"]["tables"]["people"]["columns"]["active"]["type"] == "INTEGER"


def test_csv_blank_and_duplicate_headers_are_imported_with_generated_names(setup):
    session = "6" * 32
    result = upload(session, "departments.csv", b"id,,Name,name\n1,Sales,Sales dept,Alias\n")
    columns = result["schema"]["tables"]["departments"]["columns"]
    assert list(columns) == ["id", "column_2", "Name", "name_2"]
    query = run_query("show imported columns", lambda *args: ModelQuery(
        sql='SELECT id, column_2, "Name", name_2 FROM departments', explanation="show row"), session_id=session)
    assert query["status"] == "success"
    assert query["rows"] == [{"id":1,"column_2":"Sales","Name":"Sales dept","name_2":"Alias"}]


def test_xlsx_imports_each_nonempty_sheet_and_rejects_formulas_and_merged_cells(setup):
    session = "2" * 32
    book = Workbook()
    employees = book.active
    employees.title = "Employee & Team"
    employees.append(["employee_id", "name", "department"])
    employees.append([1, "Avery", "Engineering"])
    departments = book.create_sheet("Departments")
    departments.append(["department_id", "department_name"])
    departments.append([10, "Engineering"])
    empty = book.create_sheet("Empty")
    out = io.BytesIO(); book.save(out)
    imported = upload(session, "staff.xlsx", out.getvalue())
    assert imported["source_type"] == "xlsx"
    assert imported["schema"]["source_names"] == {"Employee_Team":"Employee & Team", "Departments":"Departments"}
    assert set(imported["schema"]["tables"]) == {"Employee_Team", "Departments"}
    query = run_query("list imported staff", lambda *args: ModelQuery(sql='SELECT "name", department FROM Employee_Team WHERE employee_id = 1', explanation="staff"), session_id=session)
    assert query["status"] == "success"
    assert query["rows"] == [{"name":"Avery","department":"Engineering"}]

    formula_book = Workbook(); ws = formula_book.active; ws.append(["value"]); ws.append(["=1+1"])
    formula_file = io.BytesIO(); formula_book.save(formula_file)
    with pytest.raises(HTTPException) as error:
        upload(session, "formula.xlsx", formula_file.getvalue())
    assert "formula" in error.value.detail.lower()

    merged_book = Workbook(); ws = merged_book.active; ws.append(["name", "team"]); ws.append(["Avery", "Eng"]); ws.merge_cells("A1:B1")
    merged_file = io.BytesIO(); merged_book.save(merged_file)
    with pytest.raises(HTTPException) as error:
        upload(session, "merged.xlsx", merged_file.getvalue())
    assert "merged" in error.value.detail.lower()


@pytest.mark.parametrize("filename,content,expected", [
    ("bad.csv", b"name\n\xff\n", "utf-8"),
    ("bad.csv", b"name\nA,B\n", "more fields"),
    ("bad.json", b"{}", "choose a sqlite"),
])
def test_import_upload_validation_errors_are_clear_and_preserve_active_source(setup, filename, content, expected):
    session = "3" * 32
    active_id = upload(session, "ok.db", make_database(setup / "ok.db", "things", "safe"))["id"]
    with pytest.raises(HTTPException) as error:
        upload(session, filename, content)
    assert expected in error.value.detail.lower()
    assert main.active(request_for(session), Response())["id"] == active_id


def test_concurrent_queries_use_each_session_database(setup):
    sessions, sqls = ["f" * 32, "0" * 32], []
    for index, session in enumerate(sessions):
        source = setup / f"{index}.db"
        content = make_database(source, f"records{index}", f"db{index}")
        upload(session, source.name, content)
        sqls.append(f"SELECT name FROM parent WHERE name='db{index}'")
    def query(index):
        return run_query("find name", lambda *args: ModelQuery(sql=sqls[index], explanation="find"), session_id=sessions[index])["rows"]
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert list(pool.map(query, range(2))) == [[{"name":"db0"}], [{"name":"db1"}]]
